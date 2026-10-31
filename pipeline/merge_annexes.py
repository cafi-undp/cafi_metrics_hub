"""Merge partner indicator annexes into data/ProjectData.xlsx, the single source of results.

    python -m pipeline.merge_annexes                # dry run: prints what would change
    python -m pipeline.merge_annexes --report       # dry run, and writes docs/annex-import.md
    python -m pipeline.merge_annexes --write        # updates data/ProjectData.xlsx and the report

Re-running after a write adds nothing: merged rows are then part of the workbook.

Rules (documented in docs/annex-import.md):
1. Annex codes are mapped to standard indicators with the reference crosswalk (territorial annex)
   or by label (cross-cutting annex). Project-management indicators are not imported.
2. A column labelled "en cumulatif" (or a dated "31.12.yyyy" column) is a running total. Other year
   columns are ambiguous: the CAFI M&E guidelines ask for both the period result and the result
   since the start. A row is read as per-period results, and summed, only when its values go down
   at some point, which a running total cannot do. Otherwise it is read as a running total, the
   conservative reading. Percentages and ratios are never summed.
3. A later annex replaces an earlier one from its first reported year onward, but a zero in the
   later annex does not erase a positive value an earlier annex reported for the same year.
4. Existing workbook cells are never overwritten. Differences are listed as conflicts.
5. For a programme already in the workbook, only years after its last reported year are added.
   If the annex disagrees with the workbook on a year both report, new years are added as
   increments on top of the workbook's last value, so the series stays internally consistent.
6. A programme or indicator not in the workbook is added with its full series, if it has at least
   one non-zero value. One row per programme and indicator: when an annex reports the same
   indicator in several places, the latest file wins, then the row labelled as a total, then the
   largest value.
7. Area indicators that feed the headline hectares (FOR-01–04, AGR-01–03) are not added as new rows
   to a programme that already reports area in the workbook: partners often report the same
   hectares under several codes (wood-energy plantations inside "agroforestry + energy", set-aside
   savanna under conservation), and a new row would count them twice. Existing area rows still
   get new years. People counts above a plausibility ceiling, and rows whose unit contradicts the
   standard unit (income in $ for a % change), are also held for confirmation. The template's "ha"
   on "Nombre de bénéficiaires directs" is read as a count.
8. Partner estimates of emission reductions (CLM-01) are listed in the report but not imported:
   the hub computes carbon itself with method 2.0, and two competing figures would mislead.
Every written cell is logged in the "AnnexImport" sheet of the workbook and in the report.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
import yaml

from pipeline.annexes import read_annex

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "data" / "ProjectData.xlsx"
ANNEX_DIR = ROOT / "data" / "annexes"
REPORT = ROOT / "docs" / "annex-import.md"
FIRST_YEAR, LAST_YEAR = 2020, 2026

# --- Territorial annex (Annexe 1): framework code -> standard code. From 09_Map_Framework. ---
A1 = {
    "G.1": "GOV-03", "G.2": "GOV-08", "G.4": "GOV-06",
    "AT.1": "LUT-01", "AT.2": "LUT-02", "AT.4": "LUT-04", "AT.5": "LUT-05", "AT.8": "LUT-06", "AT.9": "LUT-07",
    "FON.1": "GOV-02", "FON.2": "LUT-08", "FON.3": "LUT-03", "FON.4": "LUT-09",
    "AG.1": "AGR-08", "AG.2": "AGR-01", "AG.4": "PPL-01",
    "E.1": "FOR-01", "E.2": "FOR-02",
    "FOR.1": "FOR-06", "FOR.3": "FOR-03", "FOR.5": "FOR-04", "FOR.7": "PPL-06", "FOR.8": "FOR-05",
}
# Sub-rows under a parent code, matched on the start of the label.
A1_SUB = [
    ("AG.3", r"^a\)", "AGR-02"), ("AG.3", r"^b\)", "AGR-03"),
    ("AT.7", r"^a\)", "PPL-03"), ("AT.7", r"^b\)", "PPL-04"),
    ("FOR.4", r"^c\.", "FOR-11"), ("FOR.4", r"^e\.", "GOV-01"),
    ("E.3", r"foyer|cuisson|cookstove", "ENR-01"),
]
# Rows without a framework code (impact/outcome sheets, cross-cutting annex): matched on the label.
BY_LABEL = [
    (r"revenus? monétaires?|augmentation des revenus|évolution des revenus", "PPL-05"),
    (r"réduction des émissions|augmentations? d.absorption", "CLM-01"),
    (r"bénéficiaires directs", "PPL-01"),
    (r"bénéficiaires indirects", "PPL-06"),
    (r"emplois créés", "PPL-02"),
    (r"cuissons? propres? propos|solutions? de cuisson.*(distribu|propos)", "ENR-01"),
    (r"taux d.adoption.*cuisson", "ENR-02"),
    (r"montant des ventes.*cuisson", "ENR-04"),
    (r"superficies? de forêts? sous plan de gestion", "FOR-06"),
    (r"extractives? forest.friendly", "FIN-09"),
    (r"années? couple de protection", "INC-01"),
    (r"productions? issues? des plantations", "AGR-04"),
    (r"revenus? agricoles?", "AGR-05"),
    (r"chaînes? de valeur sont opérationnelles", "AGR-06"),
    (r"rendements? des cultures", "AGR-07"),
    (r"alertes? de déforestation", "FOR-08"),
]
# Cross-cutting annex (Annexe 2) product codes.
A2 = {"G1": "GOV-08", "M1": "FIN-01", "R1": "GOV-02", "R2": "GOV-02", "R3": "GOV-02", "R4": "GOV-04",
      "R5": "GOV-05", "D3": "GOV-01"}
# People counts above these per-programme values are held for confirmation instead of imported.
PLAUSIBLE_MAX = {"PPL-01": 100_000, "PPL-02": 10_000, "PPL-06": 1_000_000, "GOV-01": 50_000, "GOV-08": 50_000}
NOT_SUMMED = {"LUT-10", "ENR-02", "ENR-03", "AGR-06", "AGR-07", "PPL-05", "FOR-07", "CLM-03", "FIN-07"}
UNIT_OUT = {"ha": "Hectare", "km": "Km", "USD": "USD", "%": "%", "tCO2e": "tCO2e", "Tonnes": "Tonnes"}
OUTCOME = {"GOV": "Governance", "FOR": "Forestry", "FIN": "Finance", "LUT": "Land management", "AGR": "Agriculture",
           "ENR": "Energy", "CLM": "Climate", "INC": "Inclusion", "PPL": "Infrastructures"}


def unit_class(u: str | None) -> str | None:
    """Coarse unit family, or None when the unit is blank or unrecognised (then no check is made)."""
    u = (u or "").replace("\xa0", " ").strip().lower()
    if not u:
        return None
    if re.fullmatch(r"ha|hectares?", u):
        return "ha"
    if u in ("km", "kms"):
        return "km"
    if re.fullmatch(r"\$|usd|us\$|\$ ?us|dollars?", u):
        return "usd"
    if u.startswith("%") or "pourcent" in u:
        return "%"
    if re.fullmatch(r"nb|nbre|nombre|#|personnes?|people|number|ménages?|fte.*|emplois?", u):
        return "count"
    if re.search(r"tco2|teq", u):
        return "tco2e"
    if re.fullmatch(r"t|tonnes?|t/an|kg.*", u):
        return "mass"
    return None


STANDARD_UNIT_CLASS = {"ha": "ha", "km": "km", "USD": "usd", "%": "%", "% change": "%", "People": "count",
                       "Number": "count", "FTE jobs": "count", "tCO2e": "tco2e", "Tonnes": "mass"}


def standard_code(row, family: str) -> str | None:
    label = row.label.lower()
    if row.code:
        if row.code in A1:
            return A1[row.code]
        if family == "annexe2" and row.code in A2:
            return A2[row.code]
    if row.parent_code:
        for parent, pat, code in A1_SUB:
            if row.parent_code == parent and re.search(pat, label):
                return code
        if row.parent_code in A1 or row.parent_code in A2:
            return None  # a breakdown ("dont…") of a mapped parent: the parent carries the total
    if row.code and not re.fullmatch(r"\d{1,2}", row.code):
        return None  # a framework code with no standard equivalent (project management, partial matches)
    for pat, code in BY_LABEL:
        if re.search(pat, label):
            return code
    return None


def period_values(row) -> list[tuple[int, float, bool]]:
    """(year, value, is_running_total) for each reported (non-plan) value column."""
    out = []
    for v in row.values:
        if v["plan"] or v["value"] is None or not v["years"]:
            continue
        year = max(v["years"])
        if year < FIRST_YEAR or year > LAST_YEAR:
            continue
        running = v["cumulative"] or bool(re.search(r"\d{1,2}[./]\d{1,2}[./]20\d\d", v["header"]))
        out.append((year, v["value"], running))
    return out


def series_from(rows_values, code: str) -> tuple[dict[int, float], str]:
    """Running-total series by year from one file's values for one indicator row, and how it was read."""
    vals = sorted(rows_values)
    if code in NOT_SUMMED:
        return {y: v for y, v, _ in vals}, "as reported (not summable)"
    if any(r for _, _, r in vals):
        return {y: v for y, v, _ in vals}, "running total (column marked cumulative)"
    goes_down = any(b[1] < a[1] for a, b in zip(vals, vals[1:]))
    if not goes_down:
        return {y: v for y, v, _ in vals}, "running total (values never go down)"
    total, out = 0.0, {}
    for y, v, _ in vals:
        total += v
        out[y] = total
    return out, "per-period results summed (values go down)"


def names_fallback(code: str) -> str:
    return code


def short_label(label: str) -> str:
    s = re.sub(r"^\s*[a-f][.)]\s*", "", label)
    s = re.split(r",? ventil|, désagrég|\(", s)[0]
    return s.strip(" .:-")[:70].rstrip(" ,.:;-")


def read_workbook():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["CumulResults"]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {h.lower(): i for i, h in enumerate(header)}
    years = {int(h): i for i, h in enumerate(header) if re.fullmatch(r"20\d\d", h)}
    cfg = yaml.safe_load(open(ROOT / "config" / "indicators.yml"))
    legacy = {str(k).zfill(2): v for k, v in cfg["legacy_codes"].items()}
    overrides = cfg.get("label_overrides") or {}
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, i + 1).value for i in range(len(header))]
        if not vals[col["programme"]]:
            continue
        ind = str(vals[col["indicator"]] or "").strip()
        label = str(vals[col["indicator synonym"]] or "").strip()
        m = re.match(r"^([A-Z]{3}-\d{2})", ind)
        code = m.group(1) if m else overrides.get(label) or legacy.get((re.match(r"^(\d{1,2})", ind) or [None, ""])[1].zfill(2))
        rows.append({"row": r, "programme": str(vals[col["programme"]]).strip(), "code": code, "label": label,
                     "values": {y: vals[i] for y, i in years.items()}, "country": vals[col["pays"]],
                     "partner": vals[col["implementing org."]], "start": vals[col["start year"]], "end": vals[col["end year"]]})
    return wb, ws, col, years, rows


def build(write: bool, report: bool = False) -> dict:
    cfg = yaml.safe_load(open(ROOT / "config" / "annexes.yml"))
    ref = openpyxl.load_workbook(ROOT / "config" / "reference" / "CAFI_Standard_Indicators_MetricsHub.xlsx", read_only=True, data_only=True)
    names = {r[0]: (r[1], r[13]) for r in ref["03_Indicators"].iter_rows(min_row=4, values_only=True) if r[0]}
    wb, ws, col, ycol, hub = read_workbook()
    area = set(yaml.safe_load(open(ROOT / "config" / "indicators.yml"))["area_headline"])

    # 1. Annex series per programme × standard code × annex row label.
    series: dict[tuple, dict] = {}
    unmapped = defaultdict(int)
    partner_carbon = []
    for f in sorted(cfg["files"], key=lambda f: str(f["as_of"])):
        rows = read_annex(ANNEX_DIR / f["file"], f["file"])
        for row in rows:
            code = standard_code(row, row.family)
            pv = period_values(row)
            if not code:
                if pv:
                    unmapped[(row.code or row.parent_code or "-", row.label[:60])] += 1
                continue
            if code == "CLM-01":
                if pv:
                    partner_carbon.append({"programme": f["programme"], "file": f["file"], "label": row.label[:80],
                                           "values": {y: v for y, v, _ in pv}})
                continue
            if not pv and row.targets.get("target_final") is None:
                continue
            # The annex family and sheet are part of the key: two rows can share a label but measure
            # different things (direct beneficiaries of agriculture support vs of the whole project).
            key = (f["programme"], code, short_label(row.label) or names_fallback(code), row.family, row.sheet.split("(")[0].strip())
            s = series.setdefault(key, {"values": {}, "source": {}, "target": None, "unit": row.unit, "basis": {},
                                        "as_of": "", "total": bool(re.search(r"\btotal\b", row.label, re.I))})
            ser, basis = series_from(pv, code)
            if ser:
                # A newer annex replaces the older one from its first reported year onward, except that
                # a zero does not erase a positive value already reported for that year: templates
                # that merge years into one column ("2022-2024") are often left at 0.
                start = min(ser)
                for y in [y for y in s["values"] if y >= start]:
                    if ser.get(y) == 0 and (s["values"][y] or 0) > 0:
                        ser.pop(y)
                        continue
                    for d in (s["values"], s["source"], s["basis"]):
                        d.pop(y, None)
            for y, v in ser.items():
                s["values"][y] = v
                s["source"][y] = f"{f['file']} · {row.sheet} r{row.row}"
                s["basis"][y] = basis
            if ser:
                s["as_of"] = max(s["as_of"], str(f["as_of"]))
            t = row.targets.get("target_revised") or row.targets.get("target_final")
            if t is not None:
                s["target"] = t

    # One series per programme × indicator (several annex rows can report the same indicator).
    best: dict[tuple, tuple] = {}
    dropped = []
    for key, s in series.items():
        vals = [v for v in s["values"].values() if v is not None]
        if not vals:
            continue
        rank = (s["as_of"], s["total"], max(vals))
        k2 = key[:2]
        if k2 not in best or rank > best[k2][0]:
            if k2 in best:
                dropped.append(best[k2][1])
            best[k2] = (rank, key)
        else:
            dropped.append(key)
    series = {best[k][1][:3]: series[best[k][1]] for k in best}
    dropped = [k[:3] for k in dropped]

    # 2. Compare with the workbook and decide.
    by_prog = defaultdict(list)
    for h in hub:
        by_prog[h["programme"]].append(h)
    actions, conflicts, review, matches = [], [], [], 0
    new_rows = []
    for (prog, code, label), s in sorted(series.items()):
        vals = {y: v for y, v in s["values"].items() if v is not None}
        if not vals:
            continue
        hub_rows = [h for h in by_prog.get(prog, []) if h["code"] == code]
        want, got = STANDARD_UNIT_CLASS.get(names.get(code, ("", ""))[1]), unit_class(s["unit"])
        if got == "ha" and want == "count" and re.match(r"(nombre|nb)\b", label, re.I):
            got = "count"  # the CAFI template lists "ha" as the unit of "Nombre de bénéficiaires directs" (AG.4)
        if want and got and want != got:
            review.append({"programme": prog, "code": code, "label": label, "annex": vals,
                           "reason": f"the annex unit ({s['unit'].strip()}) does not match the standard unit ({names[code][1]})"})
            continue
        cap = PLAUSIBLE_MAX.get(code)
        if cap and max(vals.values()) > cap:
            review.append({"programme": prog, "code": code, "label": label, "annex": vals,
                           "reason": f"above the plausibility ceiling of {cap:,} for one programme; confirm with the partner"})
            continue
        if not hub_rows and code in area and any(h["code"] in area for h in by_prog.get(prog, [])):
            review.append({"programme": prog, "code": code, "label": label, "annex": vals,
                           "reason": "new area row for a programme that already reports area: possible double counting",
                           "workbook": {f"{h['code']} {h['label']}": None for h in by_prog[prog] if h["code"] in area}})
            continue
        if prog not in by_prog or not hub_rows:
            if any(abs(v) > 0 for v in vals.values()):
                new_rows.append((prog, code, label, s))
                for y, v in sorted(vals.items()):
                    actions.append({"programme": prog, "code": code, "label": label, "year": y, "value": v,
                                    "action": "new programme" if prog not in by_prog else "new indicator",
                                    "source": s["source"][y], "basis": s["basis"][y]})
            continue
        if len(hub_rows) > 1:
            review.append({"programme": prog, "code": code, "label": label, "annex": vals,
                           "reason": "the workbook splits this indicator into several rows; the annex gives one total",
                           "workbook": {h["label"]: None for h in hub_rows}})
            continue
        h = hub_rows[0]
        reported = [y for y, v in h["values"].items() if v not in (None, "")]
        last = max(reported) if reported else None
        disagree = False
        for y, v in vals.items():
            hv = h["values"].get(y)
            if hv in (None, ""):
                continue
            if abs(float(hv) - v) > max(1.0, 0.01 * max(abs(v), abs(float(hv)))):
                disagree = True
                conflicts.append({"programme": prog, "code": code, "label": h["label"], "year": y,
                                  "workbook": float(hv), "annex": v, "source": s["source"][y]})
            else:
                matches += 1
        later = sorted(y for y in vals if last is None or y > last)
        prev_annex = vals.get(last) if last is not None else None
        for y in later:
            v = vals[y]
            basis = s["basis"][y]
            if disagree and last is not None and code not in NOT_SUMMED:
                base = float(h["values"][last] or 0)
                ref_prev = prev_annex if prev_annex is not None else max([vals[k] for k in vals if k < y] or [0])
                v = base + (vals[y] - ref_prev)
                basis = f"increment on workbook {last} value (annex disagrees on earlier years)"
            actions.append({"programme": prog, "code": code, "label": h["label"], "year": y, "value": v,
                            "action": "new year", "source": s["source"][y], "basis": basis, "row": h["row"]})

    # 3. Write.
    if write:
        for a in actions:
            if a["action"] == "new year":
                ws.cell(a["row"], ycol[a["year"]] + 1, round(a["value"], 2))
        meta = cfg.get("new_programmes") or {}
        template = {p: next(iter(rs)) for p, rs in by_prog.items()}
        for prog, code, label, s in new_rows:
            info = meta.get(prog) or template.get(prog) or {}
            name, unit = names.get(code, (code, ""))
            r = ws.max_row + 1
            ws.cell(r, col["pays"] + 1, info.get("country"))
            ws.cell(r, col["programme"] + 1, prog)
            ws.cell(r, col["implementing org."] + 1, info.get("partner"))
            ws.cell(r, col["outcome"] + 1, OUTCOME.get(code[:3], code[:3]))
            ws.cell(r, col["indicator"] + 1, f"{code} {name}")
            ws.cell(r, col["indicator synonym"] + 1, label)
            ws.cell(r, col["unité"] + 1, UNIT_OUT.get(unit, "Nombre"))
            ws.cell(r, col["start year"] + 1, info.get("start"))
            ws.cell(r, col["end year"] + 1, info.get("end"))
            ws.cell(r, col["final target"] + 1, s["target"])
            for y, v in s["values"].items():
                if v is not None and y in ycol:
                    ws.cell(r, ycol[y] + 1, round(v, 2))
            for a in actions:
                if a["programme"] == prog and a["code"] == code and a["label"] == label and "row" not in a:
                    a["row"] = r
        if "AnnexImport" in wb.sheetnames:
            del wb["AnnexImport"]
        log = wb.create_sheet("AnnexImport")
        log.append(["Imported on", dt.date.today().isoformat(), "by python -m pipeline.merge_annexes", "", "", "", "", "", ""])
        log.append(["Workbook row", "Programme", "Indicator", "Label", "Year", "Value", "Action", "Basis", "Source"])
        for a in actions:
            log.append([a.get("row"), a["programme"], a["code"], a["label"], a["year"], round(a["value"], 2),
                        a["action"], a["basis"], a["source"]])
        wb.save(WORKBOOK)
        for prog in meta:
            if any(p == prog for p, *_ in new_rows):
                add_programme_link(prog, meta[prog].get("mptf") or [])

    result = {"actions": actions, "conflicts": conflicts, "review": review, "matches": matches,
              "partner_carbon": partner_carbon, "duplicates": dropped,
              "unmapped": dict(unmapped), "new_rows": [(p, c, l) for p, c, l, _ in new_rows]}
    if write or report:
        write_report(result, write)
    return result


def add_programme_link(prog: str, mptf: list[int]) -> None:
    path = ROOT / "config" / "programmes.yml"
    text = path.read_text()
    if re.search(rf"^{re.escape(prog)}:", text, re.M):
        return
    path.write_text(text.rstrip("\n") + f"\n{prog}: {mptf}{' ' * max(1, 40 - len(prog) - len(str(mptf)))}# added from the indicator annexes\n")


def fmt(v):
    return "–" if v is None else f"{v:,.0f}" if abs(v) >= 10 else f"{v:,.2f}".rstrip("0").rstrip(".")


def write_report(r: dict, written: bool) -> None:
    acts = r["actions"]
    by = defaultdict(list)
    for a in acts:
        by[a["programme"]].append(a)
    lines = [
        "# Indicator annexes merged into ProjectData.xlsx",
        "",
        f"Generated by `python -m pipeline.merge_annexes{' --write' if written else ''}` on {dt.date.today():%d %B %Y}. "
        "Sources are listed in `config/annexes.yml` and kept in `data/annexes/`. The rules are in the module docstring.",
        "",
        "| | Count |",
        "|---|---:|",
        f"| Cells added | {len(acts)} |",
        f"| … in years after the workbook's last report | {sum(a['action'] == 'new year' for a in acts)} |",
        f"| … in new indicator rows of existing programmes | {sum(a['action'] == 'new indicator' for a in acts)} |",
        f"| … in new programmes | {sum(a['action'] == 'new programme' for a in acts)} |",
        f"| Values where annex and workbook agree | {r['matches']} |",
        f"| Conflicts (workbook kept) | {len(r['conflicts'])} |",
        f"| Indicators held for review | {len(r['review'])} |",
        "",
        "## Cells added, by programme",
        "",
    ]
    for prog, items in sorted(by.items()):
        lines += [f"### {prog}", "", "| Indicator | Label | Year | Value | Change | Basis | Source |", "|---|---|---:|---:|---|---|---|"]
        for a in sorted(items, key=lambda a: (a["code"], a["label"], a["year"])):
            lines.append(f"| {a['code']} | {a['label']} | {a['year']} | {fmt(a['value'])} | {a['action']} | {a['basis']} | {a['source']} |")
        lines.append("")
    lines += ["## Conflicts: the workbook value was kept", "",
              "The annex and the workbook disagree by more than 1 %. Often the annex reports per period while the "
              "workbook is cumulative, or the two use different definitions. These need a decision by the programme team.", "",
              "| Programme | Indicator | Workbook label | Year | Workbook | Annex | Source |", "|---|---|---|---:|---:|---:|---|"]
    for c in sorted(r["conflicts"], key=lambda c: (c["programme"], c["code"], c["year"])):
        lines.append(f"| {c['programme']} | {c['code']} | {c['label']} | {c['year']} | {fmt(c['workbook'])} | {fmt(c['annex'])} | {c['source']} |")
    lines += ["", "## Held for review: not merged", ""]
    for v in sorted(r["review"], key=lambda v: (v["programme"], v["code"])):
        rows = f" Workbook rows: {'; '.join(k or '(no label)' for k in v['workbook'])}." if v.get("workbook") else ""
        lines.append(f"- **{v['programme']} · {v['code']}** ({v['label']}): {v['reason']}. "
                     f"Annex: {', '.join(f'{y}: {fmt(x)}' for y, x in sorted(v['annex'].items()))}.{rows}")
    lines += ["", "## Partner estimates of emission reductions: not imported", "",
              "Shown for comparison with the hub's own carbon figures (method 2.0).", ""]
    for c in r["partner_carbon"]:
        lines.append(f"- **{c['programme']}** ({c['file']}): {', '.join(f'{y}: {fmt(v)} tCO₂e' for y, v in sorted(c['values'].items()))}")
    lines += ["", "## Same indicator reported in several annex rows: one kept", ""]
    for p, c, l in sorted(r["duplicates"]):
        lines.append(f"- {p} · {c}: not used — \"{l}\"")
    lines += ["", "## Annex rows with values but no standard indicator", "",
              "Project-management indicators (steering committees, monitoring visits) and partial matches are not imported.", ""]
    for (code, label), n in sorted(r["unmapped"].items()):
        lines.append(f"- `{code}` {label} ({n} file{'s' if n > 1 else ''})")
    REPORT.parent.mkdir(exist_ok=True)
    REPORT.write_text("\n".join(lines) + "\n")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--write", action="store_true", help="update data/ProjectData.xlsx and the report")
    ap.add_argument("--report", action="store_true", help="write docs/annex-import.md without changing the workbook")
    a = ap.parse_args()
    r = build(a.write, a.report)
    print(f"cells to add: {len(r['actions'])} | agree: {r['matches']} | conflicts: {len(r['conflicts'])} | "
          f"held for review: {len(r['review'])} | new rows: {len(r['new_rows'])}" + (f" | report: {REPORT.relative_to(ROOT)}" if a.write or a.report else ""))


if __name__ == "__main__":
    main()
