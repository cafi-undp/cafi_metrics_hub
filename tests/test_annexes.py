"""Reading partner indicator annexes and the rules used to merge them into ProjectData.xlsx."""
from pathlib import Path

import openpyxl
import pytest

from pipeline.annexes import Row, number, read_annex
from pipeline.merge_annexes import series_from, short_label, standard_code

ROOT = Path(__file__).resolve().parent.parent


@pytest.mark.parametrize("raw,expected", [
    (12, 12.0), ("1 664", 1664.0), ("3313,5", 3313.5), ("1,234,567", 1234567.0), ("30%", 30.0),
    ("N/A", None), ("n/d", None), ("-", None), ("", None), ("À déterminer en baseline", None), ("voir rapport", None),
])
def test_number(raw, expected):
    assert number(raw) == expected


def _row(code=None, parent=None, label="x", family="annexe1"):
    return Row("f", "s", 1, family, code, parent, label, None)


@pytest.mark.parametrize("code,parent,label,family,expected", [
    ("AG.2", None, "Superficies vivrières", "annexe1", "AGR-01"),
    (None, "AG.3", "a) Cultures pérennes sans déforestation", "annexe1", "AGR-02"),
    (None, "AG.3", "b) Plantations agroforestières", "annexe1", "AGR-03"),
    (None, "AG.4", "Dont - Nombre de ménages", "annexe1", None),       # breakdown of a mapped parent
    ("G.5", None, "Nombre de visites de suivi", "annexe1", None),         # project management: not imported
    ("D3", None, "Nombre de personnes formées", "annexe2", "GOV-01"),
    ("G3", None, "Nombre de réunions du COPIL", "annexe2", None),
    ("1", None, "Nombre de bénéficiaires directs, désagrégé par le genre", "annexe2", "PPL-01"),
    ("2", None, "Nombre de bénéficiaires indirects", "annexe2", "PPL-06"),
    (None, None, "(Augmentation des) Revenus monétaires directs pour les bénéficiaires directs", "annexe1", "PPL-05"),
])
def test_standard_code(code, parent, label, family, expected):
    assert standard_code(_row(code, parent, label, family), family) == expected


def test_running_total_when_column_says_cumulative():
    ser, basis = series_from([(2025, 10.0, True), (2026, 8.0, True)], "PPL-01")
    assert ser == {2025: 10.0, 2026: 8.0} and "cumulative" in basis


def test_values_that_never_go_down_are_read_as_running_total():
    # PSFD food crops: 1,664 ha in 2025 and 2026 is a stock, not 3,328 ha.
    ser, basis = series_from([(2024, 0.0, False), (2025, 1664.0, False), (2026, 1664.0, False)], "AGR-01")
    assert ser[2026] == 1664.0 and "never go down" in basis


def test_values_that_go_down_are_per_period_and_summed():
    # PIREDD Kwilu land arrangements: 176, 100, 50 per year.
    ser, basis = series_from([(2021, 176.0, False), (2022, 100.0, False), (2023, 50.0, False)], "LUT-09")
    assert ser == {2021: 176.0, 2022: 276.0, 2023: 326.0} and "summed" in basis


def test_ratios_are_never_summed():
    ser, _ = series_from([(2021, 30.0, False), (2022, 20.0, False)], "ENR-02")
    assert ser == {2021: 30.0, 2022: 20.0}


def test_short_label_drops_breakdown_text_and_trailing_punctuation():
    assert short_label("a) Cultures pérennes sans déforestation (savane ou complexe agricole)") == "Cultures pérennes sans déforestation"
    assert not short_label("Nombre de document de politique élaborés avec l’appui du projet, dont approuvés").endswith(" ")


def test_reads_a_territorial_annex(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agriculture (4+5)"
    ws.append([None, "Indicateurs de produits (4)"])                                   # banner, not the header
    ws.append([None, "#", "Indicateurs", "Unité", "Cible", None, None, "Valeur", None, None, "Commentaire"])
    ws.append([None, None, "NB : les intitulés en bleu", None, "finale (prodoc)", "mi-parcours", "2026 (PTBA)",
               "2024", "2025 (en cumulatif)", "2026 (en cumulatif)"])
    ws.append([None, "AG.2", "Superficies vivrières établies", "ha", 12000, 5000, 2237, "N/A", "1 664", 1664, "ok"])
    ws.append([None, None, "a) Cultures pérennes", None, 5250, None, None, 10, 20, 30])
    path = tmp_path / "annex.xlsx"
    wb.save(path)
    rows = read_annex(path)
    ag2 = next(r for r in rows if r.code == "AG.2")
    assert ag2.targets["target_final"] == 12000 and ag2.unit == "ha" and ag2.comment == "ok"
    reported = [(v["years"], v["value"], v["cumulative"]) for v in ag2.values if not v["plan"]]
    assert reported == [([2024], None, False), ([2025], 1664.0, True), ([2026], 1664.0, True)]
    assert [v["value"] for v in ag2.values if v["plan"]] == [2237.0]
    sub = next(r for r in rows if r.label.startswith("a)"))
    assert sub.code is None and sub.parent_code == "AG.2"


def test_import_is_recorded_in_the_workbook():
    """The committed workbook carries the provenance of every imported cell."""
    wb = openpyxl.load_workbook(ROOT / "data" / "ProjectData.xlsx", read_only=True)
    assert "AnnexImport" in wb.sheetnames
    rows = list(wb["AnnexImport"].iter_rows(min_row=3, values_only=True))
    assert rows and all(r[8] for r in rows)  # every logged cell names its source annex


@pytest.mark.parametrize("unit,expected", [
    ("ha", "ha"), ("Hectares", "ha"), ("$\xa0", "usd"), ("USD", "usd"), ("%", "%"), ("Nb", "count"),
    ("nombre", "count"), ("#", "count"), ("tCO2", "tco2e"), ("T, L", None), ("", None), (None, None),
])
def test_unit_class(unit, expected):
    from pipeline.merge_annexes import unit_class
    assert unit_class(unit) == expected
